#!/usr/bin/env python3
"""Structural validator for a generated financial model.

Catches the defects that are invisible in a spreadsheet until an input changes:
constant formulas, literals embedded in calculation rows, Excel error values,
gaps in a calculation run, undefined named ranges, merged cells written off the
top-left anchor, and sheets that are entirely static.

openpyxl cannot recalculate a workbook, so this is a structural check only:
it proves the workbook is wired like a model, not that its arithmetic is right.
Verify arithmetic by mirroring it in Python (see the parent skill).

Usage:
    python validate_model.py MODEL.xlsx [--sheet SHEET]... [--json]

Exit code 0 = no errors (warnings may exist), 1 = at least one error.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import asdict, dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

# Sheets whose contents are legitimately literals.
LITERAL_SHEETS = {
    "inputs",
    "input",
    "assumptions",
    "assumption",
    "notes",
    "note",
    "cover",
    "cover page",
    "readme",
    "checks",
    "check",
}

ERROR_LITERAL = re.compile(r"^#(REF|DIV/0|VALUE|NAME|N/A|NULL|NUM|GETTING_DATA)\b", re.I)
# A bare cell reference or range, optionally sheet-qualified.
A1_RANGE = re.compile(
    r"^(\$?[A-Za-z]{1,3}\$?\d+)?(:\$?[A-Za-z]{1,3}\$?\d+)?$|^'.+'!|^[A-Za-z_][A-Za-z0-9_ ]*!"
)
FUNC_NAMES = re.compile(r"^[A-Z][A-Z0-9_.]*$")


@dataclass
class Finding:
    level: str  # "error" | "warn" | "info"
    sheet: str
    cell: str
    message: str


def _is_literal_sheet(name: str) -> bool:
    return name.strip().lower() in LITERAL_SHEETS


def _formula_refs(formula: str) -> tuple[list[str], bool]:
    """Return (referenced ranges/names, has_numeric_literal) for a formula."""
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return [], False
    refs: list[str] = []
    numeric = False
    for tok in tokens:
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            refs.append(tok.value)
        elif tok.type == "OPERAND" and tok.subtype == "NUMBER":
            numeric = True
    return refs, numeric


def _undefined_names(ref: str, defined: set[str]) -> str | None:
    """If a reference looks like a defined name and is not one, return it."""
    raw = ref.strip()
    if "!" in raw:
        return None
    raw = raw.replace("$", "")
    if A1_RANGE.match(raw):
        return None
    if FUNC_NAMES.match(raw) and raw.upper() in {"TRUE", "FALSE"}:
        return None
    return raw if raw not in defined else None


def validate(path: Path, only: set[str] | None = None) -> list[Finding]:
    findings: list[Finding] = []

    wb = load_workbook(path, data_only=False)
    try:
        wb_values = load_workbook(path, data_only=True)
    except Exception:
        wb_values = None

    defined = set(wb.defined_names.keys())
    sheet_names = [ws.title for ws in wb.worksheets]
    if not any("check" in n.lower() for n in sheet_names):
        findings.append(
            Finding("warn", "-", "-", "no Checks sheet — model invariants are untested")
        )

    for ws in wb.worksheets:
        if only and ws.title not in only:
            continue
        literal_sheet = _is_literal_sheet(ws.title)
        formulas = 0
        literals = 0

        for row in ws.iter_rows():
            row_formulas = [(c, c.value) for c in row if isinstance(c.value, str) and c.value.startswith("=")]
            row_numeric = [
                c
                for c in row
                if isinstance(c.value, (int, float)) and not isinstance(c.value, bool)
            ]
            formulas += len(row_formulas)
            literals += len(row_numeric)

            for cell, formula in row_formulas:
                loc = f"{ws.title}!{cell.coordinate}"

                refs, numeric = _formula_refs(formula)
                if not refs and numeric:
                    findings.append(
                        Finding(
                            "error",
                            ws.title,
                            cell.coordinate,
                            f"constant formula (no cell references): {formula[:60]}",
                        )
                    )
                for ref in refs:
                    bad = _undefined_names(ref, defined)
                    if bad:
                        findings.append(
                            Finding(
                                "error",
                                ws.title,
                                cell.coordinate,
                                f"reference to undefined name '{bad}'",
                            )
                        )

            if literal_sheet or not row_formulas or len(row_formulas) < 2:
                continue

            # Numeric literal sitting inside a calculation row.
            for cell in row_numeric:
                if cell.column == 1:
                    continue  # row label
                findings.append(
                    Finding(
                        "warn",
                        ws.title,
                        cell.coordinate,
                        f"literal {cell.value!r} in a row with {len(row_formulas)} formulas "
                        "— should this be an input on the Inputs sheet, or a formula?",
                    )
                )

            # Blank gap inside a contiguous formula run.
            cols = sorted(c.column for c, _ in row_formulas)
            if len(cols) >= 4:
                for col in range(cols[0], cols[-1] + 1):
                    cell = ws.cell(row=row[0].row, column=col)
                    if cell.value is None:
                        findings.append(
                            Finding(
                                "warn",
                                ws.title,
                                cell.coordinate,
                                "blank cell inside a calculation run — a gap will "
                                "silently drop out of subtotals",
                            )
                        )

        # Cached error values, when the file has been calculated at least once.
        if wb_values is not None:
            ws_v = wb_values[ws.title]
            for row in ws_v.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and ERROR_LITERAL.match(cell.value):
                        findings.append(
                            Finding(
                                "error",
                                ws.title,
                                cell.coordinate,
                                f"Excel error value cached: {cell.value}",
                            )
                        )

        # Merged cells with values off the top-left anchor.
        for rng in ws.merged_cells.ranges:
            anchor = ws.cell(row=rng.min_row, column=rng.min_col)
            for row in ws.iter_rows(
                min_row=rng.min_row,
                max_row=rng.max_row,
                min_col=rng.min_col,
                max_col=rng.max_col,
            ):
                for cell in row:
                    if cell.coordinate == anchor.coordinate:
                        continue
                    if cell.value not in (None, ""):
                        findings.append(
                            Finding(
                                "error",
                                ws.title,
                                cell.coordinate,
                                f"value written inside merged range {rng} — invisible "
                                "in Excel; only the top-left cell holds data",
                            )
                        )

        if not literal_sheet and formulas == 0 and literals >= 20:
            findings.append(
                Finding(
                    "warn",
                    ws.title,
                    "-",
                    f"sheet has {literals} literals and no formulas — nothing here "
                    "recomputes when an input changes",
                )
            )

    return findings


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("--sheet", action="append", default=None, help="limit to sheet(s)")
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    args = ap.parse_args(argv)

    if not args.workbook.exists():
        print(f"no such workbook: {args.workbook}", file=sys.stderr)
        return 2

    findings = validate(args.workbook, set(args.sheet) if args.sheet else None)

    if args.json:
        print(json.dumps([asdict(f) for f in findings], indent=2))
    else:
        order = {"error": 0, "warn": 1, "info": 2}
        for f in sorted(findings, key=lambda f: (order[f.level], f.sheet, f.cell)):
            where = f"{f.sheet}!{f.cell}" if f.sheet != "-" else "-"
            print(f"[{f.level.upper():5}] {where}: {f.message}")
        errors = sum(1 for f in findings if f.level == "error")
        warns = sum(1 for f in findings if f.level == "warn")
        print(f"\n{errors} error(s), {warns} warning(s)")
        if errors == 0 and warns == 0:
            print("structure is clean")

    return 1 if any(f.level == "error" for f in findings) else 0


if __name__ == "__main__":
    sys.exit(main())
