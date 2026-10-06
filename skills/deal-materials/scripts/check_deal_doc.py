#!/usr/bin/env python3
"""Structural checker for deal documents (CIM, teaser, process letter, tracker notes).

Reads a .docx or .md draft and reports the defects that make a process document
unusable before it goes out: placeholder text, headings that jump a level or sit
over an empty section, financial tables whose header states no unit, numbers in
prose that do not trace back to the locked fact sheet (and the near-miss that means
one section drifted from another), and identity leaks in a blind teaser.

python-docx is used when the input is a .docx; everything else is stdlib.

Usage:
    python check_deal_doc.py DRAFT.docx --facts facts.json
    python check_deal_doc.py teaser.docx --blind --name "Atlas Holdings" \\
        --forbid "Atlas" --forbid "Fort Worth" --allow-domain advisor-bank.com
    python check_deal_doc.py tracker-notes.md --facts facts.csv --json

Exit code 0 = no errors (warnings may exist), 1 = at least one error, 2 = bad input.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

# --------------------------------------------------------------------------------------
# findings


@dataclass
class Finding:
    level: str  # "error" | "warning"
    code: str
    where: str
    message: str

    def render(self) -> str:
        return f"{self.level.upper():7s} [{self.code}] {self.where}: {self.message}"


class Report:
    def __init__(self) -> None:
        self.findings: list[Finding] = []
        self.numbers_checked = 0

    def error(self, code: str, where: str, message: str) -> None:
        self.findings.append(Finding("error", code, where, message))

    def warn(self, code: str, where: str, message: str) -> None:
        self.findings.append(Finding("warning", code, where, message))

    @property
    def errors(self) -> int:
        return sum(1 for f in self.findings if f.level == "error")


# --------------------------------------------------------------------------------------
# document model


@dataclass
class Block:
    kind: str  # "heading" | "para" | "table" | "header"
    text: str
    where: str
    level: int | None = None


@dataclass
class TableData:
    rows: list[list[str]]
    where: str


@dataclass
class Doc:
    blocks: list[Block]
    tables: list[TableData]
    props: dict[str, str]
    path: Path


def _clean(text: str) -> str:
    text = unicodedata.normalize("NFKC", text or "")
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", " ", text).strip()


def load_docx(path: Path) -> Doc:
    try:
        from docx import Document
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
    except ImportError:  # pragma: no cover
        raise SystemExit(
            "reading .docx needs python-docx (baked into the runner image)"
        ) from None

    document = Document(str(path))
    blocks: list[Block] = []
    tables: list[TableData] = []
    para_i = 0

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = _clean(para.text)
            style = (para.style.name if para.style is not None else "") or ""
            m = re.match(r"Heading (\d+)", style)
            if m:
                blocks.append(Block("heading", text, f"paragraph {para_i}", int(m.group(1))))
            elif style.startswith("Title"):
                blocks.append(Block("heading", text, f"paragraph {para_i}", 0))
            elif text:
                blocks.append(Block("para", text, f"paragraph {para_i}"))
            para_i += 1
        elif tag == "tbl":
            table = DocxTable(child, document)
            rows = [[_clean(c.text) for c in r.cells] for r in table.rows]
            where = f"table {len(tables) + 1}"
            tables.append(TableData(rows, where))
            blocks.append(Block("table", "<table>", where))

    for section in document.sections:
        for part in (
            section.header,
            section.footer,
            section.first_page_header,
            section.first_page_footer,
        ):
            for para in part.paragraphs:
                text = _clean(para.text)
                if text:
                    blocks.append(Block("header", text, "header/footer"))

    core = document.core_properties
    props = {
        "author": core.author or "",
        "last_modified_by": core.last_modified_by or "",
        "title": core.title or "",
        "comments": core.comments or "",
    }
    return Doc(blocks, tables, props, path)


MD_TABLE_ROW = re.compile(r"^\s*\|.*\|\s*$")
MD_SEP = re.compile(r"^\s*\|[\s:|-]+\|\s*$")


def load_md(path: Path) -> Doc:
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    blocks: list[Block] = []
    tables: list[TableData] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        if MD_TABLE_ROW.match(line):
            buf = []
            while i < len(lines) and MD_TABLE_ROW.match(lines[i]):
                buf.append(lines[i])
                i += 1
            rows = [
                [c.strip() for c in r.strip().strip("|").split("|")]
                for r in buf
                if not MD_SEP.match(r)
            ]
            where = f"table {len(tables) + 1} (line {i - len(rows) + 1})"
            tables.append(TableData(rows, where))
            blocks.append(Block("table", "<table>", where))
            continue
        text = _clean(line)
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            blocks.append(
                Block(
                    "heading",
                    text.lstrip("# ").strip(),
                    f"line {i + 1}",
                    len(m.group(1)),
                )
            )
        elif text:
            blocks.append(Block("para", text, f"line {i + 1}"))
        i += 1
    return Doc(blocks, tables, {}, path)


def load_doc(path: Path) -> Doc:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return load_docx(path)
    if suffix in {".md", ".markdown", ".txt"}:
        return load_md(path)
    raise SystemExit(f"unsupported input {suffix}: pass a .docx or .md draft")


# --------------------------------------------------------------------------------------
# numbers


NUM_RE = re.compile(
    r"""(?<![\w.,%])
        (?P<cur>US\$|\$|USD|EUR|GBP|EUR|€|£)?
        (?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+(?:\.\d+)?)
        (?P<sp>\s*)
        (?P<mult>billion|million|thousand|bn|mm|mn|m|k|%)?
        (?P<x>x)?
    """,
    re.X | re.I,
)

SCALES = {
    "billion": 1e9,
    "bn": 1e9,
    "million": 1e6,
    "mm": 1e6,
    "mn": 1e6,
    "m": 1e6,
    "thousand": 1e3,
    "k": 1e3,
}
FACTORS = (1e-6, 1e-3, 1e3, 1e6, 1e9)
REL_EXACT = 1e-3
REL_NEAR = 0.05


@dataclass
class Number:
    text: str
    kind: str  # "money" | "pct" | "mult" | "bare"
    value: float
    negative: bool = False


def parse_number(match: re.Match[str]) -> Number | None:
    raw = match.group("num").replace(",", "")
    try:
        value = float(raw)
    except ValueError:  # pragma: no cover
        return None
    cur, mult, gap, x = (
        match.group("cur"),
        match.group("mult"),
        match.group("sp"),
        match.group("x"),
    )
    if mult:
        mult = mult.lower()
    if mult in {"m", "k"} and gap:
        return None  # "42 m" of cable is not $42M — ambiguous, do not guess
    if mult == "%":
        return Number(match.group(0).strip(), "pct", value)
    if x:
        return Number(match.group(0).strip(), "mult", value)
    if mult:
        return Number(match.group(0).strip(), "money", value * SCALES[mult])
    if cur:
        return Number(match.group(0).strip(), "money", value)
    return Number(match.group(0).strip(), "bare", value)


def scan_numbers(text: str) -> list[tuple[Number, int]]:
    out: list[tuple[Number, int]] = []
    for m in NUM_RE.finditer(text):
        n = parse_number(m)
        if n is None:
            continue
        before = text[max(0, m.start() - 1) : m.start()]
        after = text[m.end() : m.end() + 1]
        if before == "(" and after == ")":
            n.negative = True
        out.append((n, m.start()))
    return out


def is_year(n: Number) -> bool:
    return (
        n.kind == "bare" and not n.negative and n.value.is_integer() and 1900 <= n.value <= 2100
    )


def is_trivial(n: Number) -> bool:
    """Counts and small integers carry no drift risk; years are excluded outright."""
    return n.kind == "bare" and abs(n.value) < 1000


def unit_spec(unit: str) -> tuple[str, float]:
    """Map a stated unit to (kind, multiplier that converts the value to base units)."""
    u = (unit or "").strip().lower()
    if "%" in u or "percent" in u:
        return "pct", 1.0
    if re.search(r"\bx\b|multiple", u):
        return "mult", 1.0
    if re.search(r"\b(employee|customer|unit|store|site|headcount|count|volume)", u):
        return "count", 1.0
    if "billion" in u or re.search(r"\bbn\b", u):
        return "money", 1e9
    if "million" in u or re.search(r"\b(?:mm|mn|m)\b", u):
        return "money", 1e6
    if "thousand" in u or "000" in u:
        return "money", 1e3
    return "money", 1.0


@dataclass
class Fact:
    fid: str
    label: str
    kind: str
    raw: float
    base: float
    unit: str

    @property
    def candidates(self) -> tuple[float, ...]:
        return (self.raw, self.base)


def load_facts(path: Path) -> list[Fact]:
    records: list[dict] = []
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if isinstance(data, dict) and isinstance(data.get("facts"), list):
            records = data["facts"]
        elif isinstance(data, dict):
            records = [{"id": k, "value": v} for k, v in data.items()]
        elif isinstance(data, list):
            records = data
    elif path.suffix.lower() == ".csv":
        with path.open(newline="", encoding="utf-8") as fh:
            for row in csv.DictReader(fh):
                records.append({k.lower().strip(): v for k, v in row.items() if k})
    else:
        raise SystemExit("--facts takes a .json or .csv sheet")

    facts: list[Fact] = []
    for rec in records:
        if not isinstance(rec, dict):
            continue
        label = str(
            rec.get("label") or rec.get("metric") or rec.get("name") or rec.get("id") or "?"
        )
        fid = str(rec.get("id") or rec.get("metric") or label)
        unit = str(rec.get("unit") or "")
        raw = rec.get("value")
        if raw is None:
            continue
        if isinstance(raw, str):
            tokens = scan_numbers(raw)
            if not tokens:
                continue
            raw = tokens[0][0].value * (-1 if tokens[0][0].negative else 1)
            unit = unit or raw_unit_from_string(str(rec.get("value")))
        try:
            raw = float(raw)
        except (TypeError, ValueError):
            continue
        kind, scale = unit_spec(unit)
        facts.append(Fact(fid, label, kind, raw, raw * scale, unit))
    return facts


def raw_unit_from_string(text: str) -> str:
    u = text.lower()
    for token in ("billion", "million", "thousand", "bn", "mm", "bn", "%", "x"):
        if re.search(rf"\b{re.escape(token)}\b|{re.escape(token)}", u):
            return u
    return ""


KIND_COMPAT = {
    "money": {"money"},
    "pct": {"pct"},
    "mult": {"mult"},
    "bare": {"money", "pct", "mult", "count"},
    "count": {"count", "bare"},
}


def match_fact(n: Number, facts: list[Fact]) -> tuple[str, Fact | None]:
    """Return (verdict, fact): exact | scale | near | none."""
    value = -n.value if n.negative else n.value
    allowed = KIND_COMPAT.get(n.kind, set())
    near: Fact | None = None
    near_gap = 1.0
    for fact in facts:
        if n.kind != "bare" and fact.kind not in allowed:
            continue
        for cand in fact.candidates:
            if cand == 0:
                continue
            rel = abs(value - cand) / abs(cand)
            if rel <= REL_EXACT:
                return "exact", fact
            if rel <= REL_NEAR and rel < near_gap:
                near, near_gap = fact, rel
        if n.kind == "bare" and value:
            for factor in FACTORS:
                scaled = value * factor
                if fact.raw and abs(scaled - fact.raw) / abs(fact.raw) <= REL_EXACT:
                    return "scale", fact
                if fact.base and abs(scaled - fact.base) / abs(fact.base) <= REL_EXACT:
                    return "scale", fact
    if near is not None:
        return "near", near
    return "none", None


# --------------------------------------------------------------------------------------
# checks

PLACEHOLDERS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"\bTBD\b", re.I), "TBD"),
    (re.compile(r"\bTBC\b", re.I), "TBC"),
    (re.compile(r"\bTODO\b", re.I), "TODO"),
    (re.compile(r"\bFIXME\b", re.I), "FIXME"),
    (re.compile(r"(?<![A-Za-z])X{2,}(?![A-Za-z])"), "X placeholder"),
    (re.compile(r"\[\s*(?:insert|add|tbd|todo|xx)", re.I), "unfilled [insert …]"),
    (re.compile(r"\[\s*\]"), "empty [ ]"),
    (re.compile(r"\{\{|\}\}"), "unrendered {{template}}"),
    (re.compile(r"<<"), "<< placeholder"),
    (re.compile(r"\bLorem ipsum\b", re.I), "Lorem ipsum"),
    (
        re.compile(r"\[(?:Company|Target|Buyer|Seller|Client)\b", re.I),
        "bracketed placeholder",
    ),
]

UNIT_CUE = re.compile(
    r"\$|USD|EUR|GBP|€|£|%|percent|\b(?:x|multiple)\b|\(?\s*000\s*\)?|thousand|million|"
    r"billion|\bmm\b|\bmn\b|\bbn\b|per\s+\w+|count|#|headcount|employees|customers",
    re.I,
)
PERIOD_CUE = re.compile(
    r"\b(?:FY|CY|Q[1-4]|H[12])\s?\d{2,4}(?![0-9])|"
    r"\b(?:19|20)\d{2}[AEP]?(?![0-9A-Za-z])|"
    r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[a-z]*\b",
    re.I,
)
NUMERIC_CELL = re.compile(r"^\(?\s*[\$€£]?\s*\d[\d,]*(?:\.\d+)?\s*%?\s*\)?$")

LEAK_URL = re.compile(r"\b(?:https?://|www\.)\S+", re.I)
LEAK_EMAIL = re.compile(r"\b[\w.+-]+@([\w-]+(?:\.[\w-]+)+)\b")


def check_placeholders(doc: Doc, rep: Report) -> None:
    for block in doc.blocks:
        if block.kind == "table":
            continue
        for pattern, label in PLACEHOLDERS:
            m = pattern.search(block.text)
            if m:
                rep.error(
                    "PLACEHOLDER",
                    block.where,
                    f"unresolved {label!r} ({m.group(0)!r}) — fill it or delete it",
                )
    for table in doc.tables:
        for r, row in enumerate(table.rows, start=1):
            for c, cell in enumerate(row, start=1):
                for pattern, label in PLACEHOLDERS:
                    if pattern.search(cell):
                        rep.error(
                            "PLACEHOLDER",
                            f"{table.where} r{r}c{c}",
                            f"unresolved {label!r} in a table cell",
                        )


def check_headings(doc: Doc, rep: Report) -> None:
    items = [b for b in doc.blocks if b.kind in {"heading", "para", "table"} and b.text]
    heading_idx = [i for i, b in enumerate(items) if b.kind == "heading"]
    if not heading_idx:
        rep.warn(
            "NO_HEADINGS",
            doc.path.name,
            "no styled headings found — the document has no outline and no TOC",
        )
        return
    first = items[heading_idx[0]]
    if first.level not in (0, 1):
        rep.warn(
            "HEADING_START",
            first.where,
            f"first heading is level {first.level}; documents start at level 1",
        )
    for pos, idx in enumerate(heading_idx):
        h = items[idx]
        if pos + 1 >= len(heading_idx):
            continue
        nxt = items[heading_idx[pos + 1]]
        jump = (h.level or 0) + 1 < (nxt.level or 0)
        if jump:
            rep.error(
                "HEADING_JUMP",
                nxt.where,
                f"heading {nxt.text!r} is level {nxt.level} directly "
                f"after level {h.level} — a level is missing",
            )
        content = [b for b in items[idx + 1 : heading_idx[pos + 1]] if b.kind != "heading"]
        # a document title (or the md H1) directly above its first section is fine
        title_then_section = pos == 0 and (nxt.level or 0) == (h.level or 0) + 1
        if not content and not title_then_section and not jump:
            rep.warn(
                "ORPHAN_HEADING",
                h.where,
                f"heading {h.text!r} is followed immediately by {nxt.text!r}"
                " with nothing between — an empty section",
            )


def check_tables(doc: Doc, rep: Report, facts: list[Fact]) -> None:
    for table in doc.tables:
        if len(table.rows) < 2 or not table.rows[0]:
            continue
        header = table.rows[0]
        numeric_cols = []
        for col in range(len(header)):
            values = [row[col] for row in table.rows[1:] if col < len(row)]
            if sum(1 for v in values if NUMERIC_CELL.match(v.strip())) >= 1:
                numeric_cols.append(col)
        if not numeric_cols:
            continue
        header_has_unit = any(UNIT_CUE.search(h or "") for h in header)
        if not header_has_unit:
            rep.error(
                "TABLE_UNITS",
                table.where,
                "numeric column but the header row states no unit — put "
                "'$M unless noted' (or the real unit) in the header row",
            )
            continue
        if len(numeric_cols) >= 3 and not any(PERIOD_CUE.search(h or "") for h in header):
            rep.warn(
                "TABLE_PERIODS",
                table.where,
                "multi-column numeric table with no period in the header "
                "(FY2024A / Q3 2026E) — a reader cannot date the columns",
            )
        if not facts:
            continue
        default_header = next((h for h in header if UNIT_CUE.search(h or "")), header[0])
        default_kind, default_scale = unit_spec(default_header)
        for col in numeric_cols:
            col_unit = header[col] if col < len(header) else ""
            if UNIT_CUE.search(col_unit or ""):
                kind, scale = unit_spec(col_unit)
            else:
                kind, scale = default_kind, default_scale
            for r, row in enumerate(table.rows[1:], start=2):
                cell = (row[col] if col < len(row) else "").strip()
                if not NUMERIC_CELL.match(cell):
                    continue
                tokens = scan_numbers(cell)
                if not tokens:
                    continue
                n = tokens[0][0]
                applied = n
                rescaled = False
                if kind == "pct" and n.kind == "bare":
                    applied = Number(n.text, "pct", n.value, n.negative)
                elif kind == "money" and n.kind == "bare" and scale != 1.0:
                    applied = Number(n.text, "money", n.value * scale, n.negative)
                    rescaled = True
                verdict, fact = match_fact(applied, facts)
                rep.numbers_checked += 1
                where = f"{table.where} r{r}c{col + 1}"
                if verdict == "scale" and fact is not None:
                    rep.warn(
                        "TABLE_SCALE",
                        where,
                        f"{cell!r} matches {fact.label} only at a different scale — "
                        "confirm the column's units",
                    )
                elif verdict == "near" and fact is not None:
                    rep.warn(
                        "TABLE_DRIFT",
                        where,
                        f"{cell!r} is close to but not equal to {fact.label} — "
                        "one of the two is stale",
                    )
                elif rescaled and verdict == "none":
                    raw_verdict, raw_fact = match_fact(n, facts)
                    if raw_verdict == "exact" and raw_fact is not None:
                        rep.warn(
                            "TABLE_SCALE",
                            where,
                            f"column is headed {default_header!r} / {col_unit!r} but "
                            f"{cell!r} matches {raw_fact.label} in the sheet's own "
                            f"units ({raw_fact.unit}) — the table is off by a scale",
                        )


def check_numbers(
    doc: Doc, rep: Report, facts: list[Fact], allow: list[re.Pattern[str]]
) -> None:
    if not facts:
        rep.warn(
            "NO_FACTS",
            doc.path.name,
            "no --facts sheet: prose numbers were not traced to a source",
        )
        return
    for block in doc.blocks:
        if block.kind != "para":
            continue
        for n, _pos in scan_numbers(block.text):
            if is_year(n) or is_trivial(n):
                continue
            if any(p.search(n.text) for p in allow):
                continue
            rep.numbers_checked += 1
            verdict, fact = match_fact(n, facts)
            if verdict == "exact":
                continue
            if verdict == "scale" and fact is not None:
                rep.warn(
                    "PROSE_SCALE",
                    block.where,
                    f"{n.text!r} matches {fact.label} only at a different scale — "
                    "state the unit in the prose or fix the figure",
                )
            elif verdict == "near" and fact is not None:
                rep.error(
                    "DRIFT",
                    block.where,
                    f"{n.text!r} is not {fact.label} ({fact.raw:g} {fact.unit}) but is "
                    "within 5% of it — the prose and the fact sheet disagree",
                )
            else:
                rep.error(
                    "UNSOURCED",
                    block.where,
                    f"{n.text!r} in prose does not trace to the fact sheet — add it "
                    "to facts.json or fix the figure",
                )


def check_leaks(
    doc: Doc, rep: Report, blind: bool, names: list[str], allow_domains: list[str]
) -> None:
    if not names and not blind:
        return
    haystacks = [(b.where, b.text) for b in doc.blocks if b.kind != "table"]
    for table in doc.tables:
        for r, row in enumerate(table.rows, start=1):
            for cell in row:
                if cell:
                    haystacks.append((f"{table.where} r{r}", cell))
    for key, value in doc.props.items():
        if value:
            haystacks.append((f"document property {key}", value))
    haystacks.append(("file name", doc.path.name))

    for term in names:
        if not term:
            continue
        pat = re.compile(re.escape(term), re.I)
        for where, text in haystacks:
            if pat.search(text):
                rep.error(
                    "IDENTITY_LEAK",
                    where,
                    f"{term!r} appears in a blind document — remove it",
                )
    if not blind:
        return
    allow = [d.lower().lstrip("@") for d in allow_domains]
    for where, text in haystacks:
        for m in LEAK_URL.finditer(text):
            host = re.sub(r"^https?://", "", m.group(0)).split("/")[0].lower()
            if not any(host == d or host.endswith("." + d) for d in allow):
                rep.error(
                    "IDENTITY_LEAK",
                    where,
                    f"URL {m.group(0)!r} in a blind teaser — allow the advisor's own "
                    "domain with --allow-domain or remove it",
                )
        for m in LEAK_EMAIL.finditer(text):
            host = m.group(1).lower()
            if not any(host == d or host.endswith("." + d) for d in allow):
                rep.error(
                    "IDENTITY_LEAK",
                    where,
                    f"email {m.group(0)!r} is not on an allowed domain — the teaser "
                    "contact should be the advisor, not the company",
                )


# --------------------------------------------------------------------------------------
# main


def run(args: argparse.Namespace) -> int:
    path = Path(args.draft)
    if not path.exists():
        print(f"no file at {path}", file=sys.stderr)
        return 2
    doc = load_doc(path)
    facts = load_facts(Path(args.facts)) if args.facts else []
    allow = [re.compile(p) for p in args.allow or []]
    names = list(args.name or []) + list(args.forbid or [])

    rep = Report()
    check_placeholders(doc, rep)
    check_headings(doc, rep)
    check_tables(doc, rep, facts)
    check_numbers(doc, rep, facts, allow)
    check_leaks(doc, rep, args.blind, names, args.allow_domain or [])

    if args.json:
        print(
            json.dumps(
                {
                    "draft": str(path),
                    "errors": rep.errors,
                    "warnings": len(rep.findings) - rep.errors,
                    "numbers_checked": rep.numbers_checked,
                    "findings": [f.__dict__ for f in rep.findings],
                },
                indent=2,
            )
        )
    else:
        shown = (
            rep.findings if not args.quiet else [f for f in rep.findings if f.level == "error"]
        )
        for finding in shown:
            print(finding.render())
        if not shown:
            print("OK      no findings")
        print(
            f"\n{rep.errors} error(s), {len(rep.findings) - rep.errors} warning(s), "
            f"{rep.numbers_checked} number(s) checked against "
            f"{len(facts)} fact(s) in {path.name}"
        )
    return 1 if rep.errors else 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(
        prog="check_deal_doc.py",
        description="Structural check for a CIM, teaser, letter or tracker draft.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="Exit 0 = clean, 1 = errors found, 2 = bad input.",
    )
    parser.add_argument("draft", help="the .docx or .md draft to check")
    parser.add_argument("--facts", help="locked fact sheet: .json ({'facts': [...]}) or .csv")
    parser.add_argument(
        "--blind",
        action="store_true",
        help="blind-document mode: flag URLs and emails off the allowed domains",
    )
    parser.add_argument(
        "--name",
        action="append",
        metavar="TEXT",
        help="a term that must not appear (the target's name); repeatable",
    )
    parser.add_argument(
        "--forbid",
        action="append",
        metavar="TEXT",
        help="any other term to forbid (brand, city, customer); repeatable",
    )
    parser.add_argument(
        "--allow-domain",
        action="append",
        metavar="DOMAIN",
        help="domain allowed in a blind document (the advisor's); repeatable",
    )
    parser.add_argument(
        "--allow",
        action="append",
        metavar="REGEX",
        help="number text the fact check should ignore; repeatable",
    )
    parser.add_argument("--json", action="store_true", help="machine-readable output")
    parser.add_argument("--quiet", action="store_true", help="print errors only")
    args = parser.parse_args(argv)
    return run(args)


if __name__ == "__main__":
    sys.exit(main())
